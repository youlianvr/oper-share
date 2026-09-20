#!/usr/bin/env python3
"""extract_docx.py - extract clean text from .docx to stdout."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
def main():
    path = sys.argv[1] if len(sys.argv) > 1 else "file.docx"
    doc = Document(path)
    for p in doc.paragraphs:
        if p.text.strip(): print(p.text)
    for t in doc.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            print(" | ".join(cells))
if __name__ == "__main__":
    main()
